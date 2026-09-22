#!/usr/bin/env python3
"""
One-time (or occasional) converter: Excel source document -> indicators.json.

Reproduces the parsing rules from the former Django load_data() view.
After conversion, docs/data/indicators.json is the source of truth.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

DEFAULT_SOURCE_LINKS = {
    "BCtA": "https://www.businesscalltoaction.org/sites/default/files/BCtA-Indicator-Guidance-Note.pdf",
    "SASB": "https://www.sasb.org/standards-overview/",
    "GRI": "https://www.globalreporting.org/standards",
    "IRIS+": "https://iris.thegiin.org/standards/",
}

REQUIRED_HEADERS = [
    "SPSD Category",
    "SPSD Program Area",
    "Foreign Assistance Standard Indicators",
    "Engagement Objective",
    "Business Indicator",
    "Source",
    "Source Reference",
    "SDG Goal",
    "SDG Target",
]

INDICATOR_KEYS = [
    "spsd_category_code",
    "spsd_category",
    "program_area",
    "fa_indicator_code",
    "fa_indicator",
    "engagement_objective",
    "business_code",
    "business_indicator",
    "source",
    "source_reference",
    "sdg_goal",
    "sdg_goal_title",
    "sdg_target_number",
    "sdg_target",
]


def parse_category(raw: str) -> tuple[str | None, str]:
    category = raw.strip()
    code = None
    if ":" in category:
        code = category.split(":", 1)[0].strip()
        category = category.split(":", 1)[1].strip()
    return code, category


def parse_fa_indicator(raw: str) -> tuple[str, str]:
    """Split on first ':', else first space (matches Django load_data)."""
    text = raw.strip()
    if ":" in text:
        split_index = text.index(":")
    else:
        split_index = text.index(" ")
    code = text[:split_index].strip()
    indicator = text[split_index + 1 :].strip()
    return code, indicator


def parse_business_indicator(raw: str) -> tuple[str, str]:
    """Match Django load_data business-indicator parsing, including quirks.

    - With a colon: code is text before the first colon (stripped); indicator is
      everything after that colon (via replace(code + ':', ...).strip()).
    - Without a colon: code is the first two words; indicator stays the full
      original string (Django's replace(code + ':', ...) is a no-op).
    """
    text = str(raw)
    if ":" in text:
        code_raw = text.split(":", 1)[0]
        code = code_raw.strip()
        indicator = text.replace(code_raw + ":", "", 1).strip()
        return code, indicator
    parts = text.split(" ")
    code = (parts[0] + " " + parts[1]).strip()
    # Intentionally keep the full string — matches Django replace(code+':', ...)
    indicator = text.replace(code + ":", "").strip()
    return code, indicator


def parse_sdg_goal(raw: str | None) -> tuple[int | None, str | None]:
    if raw is None:
        return None, None
    text = str(raw).strip()
    if not text or "." not in text:
        return None, None
    # Django used split('.')[0] and split('.')[1] — title may contain periods
    parts = text.split(".", 1)
    try:
        goal = int(parts[0].strip())
    except ValueError:
        return None, None
    title = parts[1].strip()
    return goal, title


def parse_sdg_target(raw: str | None) -> tuple[str | None, str | None]:
    if raw is None:
        return None, None
    text = str(raw).strip()
    if not text:
        return None, None
    if ":" in text:
        left, right = text.split(":", 1)
        title = right.strip()
        # Django: last token of the left side is the target number
        tokens = left.strip().split()
        number = tokens[-1].strip() if tokens else None
        return number, title
    # No colon — try to extract a leading X.Y number, rest is title
    match = re.match(r"^(\d+(?:\.\d+[a-zA-Z]?)?)\s+(.*)$", text)
    if match:
        return match.group(1), match.group(2).strip()
    return None, None


def row_to_dict(header: list, cells: list) -> dict:
    return {header[i]: cells[i] for i in range(len(header))}


def convert(xlsx_path: Path, worksheet_number: int = 1, header_row: int = 1) -> dict:
    wb = load_workbook(filename=str(xlsx_path), data_only=True)
    ws = wb.worksheets[worksheet_number - 1]

    header_cells = list(ws[header_row])
    keys = [c.value for c in header_cells]

    missing = [h for h in REQUIRED_HEADERS if h not in keys]
    if missing:
        raise SystemExit(f"Missing required columns: {missing}")

    indicators = []
    warnings = []
    last_goal = None
    last_goal_title = None
    last_target_number = None
    last_target = None

    for excel_row_num, row in enumerate(
        ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1
    ):
        row_obj = row_to_dict(keys, list(row))
        if row_obj.get("SPSD Category") is None:
            break

        cat_code, cat_name = parse_category(str(row_obj["SPSD Category"]))
        fa_code, fa_text = parse_fa_indicator(str(row_obj["Foreign Assistance Standard Indicators"]))
        biz_code, biz_text = parse_business_indicator(str(row_obj["Business Indicator"]))

        engagement = row_obj.get("Engagement Objective")
        if engagement:
            engagement = str(engagement).strip()
            # Django split on ':' for nested objectives; current data has none
            if ":" in engagement:
                engagement = ": ".join(p.strip() for p in engagement.split(":"))

        source = str(row_obj["Source"]).strip()
        source_ref = row_obj.get("Source Reference")
        if source_ref is not None:
            source_ref = str(source_ref).strip() or None

        goal, goal_title = parse_sdg_goal(row_obj.get("SDG Goal"))
        target_number, target_title = parse_sdg_target(row_obj.get("SDG Target"))

        carried = []
        if goal is None:
            if last_goal is not None:
                goal, goal_title = last_goal, last_goal_title
                carried.append("sdg_goal")
            else:
                warnings.append(f"Row {excel_row_num}: no SDG Goal and nothing to carry forward")
        if target_number is None or target_title is None:
            if last_target_number is not None:
                target_number, target_title = last_target_number, last_target
                carried.append("sdg_target")
            else:
                warnings.append(f"Row {excel_row_num}: no SDG Target and nothing to carry forward")

        if carried:
            warnings.append(
                f"Row {excel_row_num}: carried forward {', '.join(carried)} from previous row"
            )

        # Track for next-row carry-forward (only when we have valid values)
        if goal is not None:
            last_goal, last_goal_title = goal, goal_title
        if target_number is not None and target_title is not None:
            last_target_number, last_target = target_number, target_title

        record = {
            "spsd_category_code": cat_code,
            "spsd_category": cat_name,
            "program_area": str(row_obj["SPSD Program Area"]).strip(),
            "fa_indicator_code": fa_code,
            "fa_indicator": fa_text,
            "engagement_objective": engagement,
            "business_code": biz_code,
            "business_indicator": biz_text,
            "source": source,
            "source_reference": source_ref,
            "sdg_goal": goal,
            "sdg_goal_title": goal_title,
            "sdg_target_number": target_number,
            "sdg_target": target_title,
        }
        # Stable key order
        indicators.append({k: record[k] for k in INDICATOR_KEYS})

    return {
        "default_source_links": DEFAULT_SOURCE_LINKS,
        "indicators": indicators,
        "_warnings": warnings,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Convert HI Excel source to indicators.json")
    parser.add_argument(
        "xlsx",
        nargs="?",
        default="data/HI_Source_Document_100621_1_pn5VMxa.xlsx",
        help="Path to source Excel file",
    )
    parser.add_argument(
        "-o",
        "--output",
        default="docs/data/indicators.json",
        help="Output JSON path",
    )
    parser.add_argument("--worksheet", type=int, default=1)
    parser.add_argument("--header-row", type=int, default=1)
    args = parser.parse_args()

    root = Path(__file__).resolve().parent.parent
    xlsx_path = Path(args.xlsx)
    if not xlsx_path.is_absolute():
        xlsx_path = root / xlsx_path
    out_path = Path(args.output)
    if not out_path.is_absolute():
        out_path = root / out_path

    result = convert(xlsx_path, worksheet_number=args.worksheet, header_row=args.header_row)
    warnings = result.pop("_warnings")
    for w in warnings:
        print(f"WARNING: {w}", file=sys.stderr)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    with out_path.open("w", encoding="utf-8") as f:
        json.dump(result, f, indent=2, ensure_ascii=False)
        f.write("\n")

    print(f"Wrote {len(result['indicators'])} indicators to {out_path}")
    if warnings:
        print(f"({len(warnings)} warnings — see stderr)", file=sys.stderr)


if __name__ == "__main__":
    main()
